from datetime import date

from PySide6.QtCore import QDate
from PySide6.QtWidgets import (
    QFileDialog,
    QMessageBox,
    QWidget,
)

from app.models.estoque_table_model import EstoqueTableModel
from app.services.estoque_service import EstoqueService
from app.utils.logger import get_logger
from app.utils.table_utils import configurar_tabela
from app.views.ui_estoque import Ui_Estoque_Tab

logger = get_logger("estoque_controller")


class EstoqueController(QWidget):
    def __init__(self):
        super().__init__()

        self.ui = Ui_Estoque_Tab()
        self.ui.setupUi(self)

        self.estoque_service = EstoqueService()
        self.model = EstoqueTableModel()

        self._configurar_tabela()
        self._configurar_campos()
        self._conectar_sinais()
        self.atualizar()

        logger.info("Tela de estoque inicializada")

    def _configurar_tabela(self):
        self.ui.tb_Estoque.setModel(self.model)
        configurar_tabela(self.ui.tb_Estoque, coluna_stretch=1, ordenavel=True)

    def _configurar_campos(self):
        self.ui.dt_Filtro.setCalendarPopup(True)
        self.ui.dt_Filtro.setDisplayFormat("dd/MM/yyyy")
        self.ui.dt_Filtro.setDate(QDate.currentDate())

    def _conectar_sinais(self):
        self.ui.bt_Atualizar.clicked.connect(self.atualizar)
        self.ui.bt_Exportar.clicked.connect(self.exportar_excel)
        self.ui.ck_Ocultar.stateChanged.connect(self._aplicar_filtro_local)

    def atualizar(self):
        """Recarrega o saldo do banco para a data selecionada."""
        data_limite = self.ui.dt_Filtro.date().toPython()
        try:
            # Busca sempre COM os zerados; ocultar/mostrar depois é só
            # um filtro em memória (não precisa ir ao banco de novo
            # toda vez que o checkbox é marcado/desmarcado).
            self._todas_as_linhas = self.estoque_service.listar_saldo(
                data_limite=data_limite, ocultar_zerados=False
            )
            self._aplicar_filtro_local()
            logger.info(
                f"Estoque atualizado | data={data_limite} | "
                f"produtos={len(self._todas_as_linhas)}"
            )
        except Exception as e:
            logger.error(f"Erro ao atualizar estoque: {e}", exc_info=True)
            QMessageBox.critical(self, "Erro", f"Erro ao atualizar: {e}")

    def _aplicar_filtro_local(self):
        linhas = getattr(self, "_todas_as_linhas", [])
        if self.ui.ck_Ocultar.isChecked():
            linhas = [
                linha for linha in linhas
                if abs(linha.get("saldo", 0)) > 1e-9
            ]
        self.model.atualizar_dados(linhas)

    def exportar_excel(self):
        linhas = self.model.obter_todos()
        if not linhas:
            QMessageBox.warning(
                self, "Aviso", "Não há dados para exportar."
            )
            return

        data_limite = self.ui.dt_Filtro.date().toPython()
        nome_sugerido = (
            f"estoque_{data_limite.strftime('%Y-%m-%d')}.xlsx"
        )

        caminho, _ = QFileDialog.getSaveFileName(
            self,
            "Exportar estoque para Excel",
            nome_sugerido,
            "Planilha Excel (*.xlsx)",
        )
        if not caminho:
            return

        try:
            self._gerar_planilha(caminho, linhas, data_limite)
            logger.info(f"Estoque exportado para Excel: {caminho}")
            QMessageBox.information(
                self, "Sucesso",
                f"Estoque exportado com sucesso para:\n{caminho}"
            )
        except Exception as e:
            logger.error(
                f"Erro ao exportar estoque para Excel: {e}", exc_info=True
            )
            QMessageBox.critical(
                self, "Erro", f"Erro ao exportar: {e}"
            )

    @staticmethod
    def _gerar_planilha(caminho, linhas, data_limite):
        from openpyxl import Workbook
        from openpyxl.styles import Font

        wb = Workbook()
        ws = wb.active
        ws.title = "Estoque"

        ws.append([
            f"Saldo de estoque até {data_limite.strftime('%d/%m/%Y')}"
        ])
        ws["A1"].font = Font(bold=True, size=12)
        ws.append([])  # linha em branco

        cabecalho = [
            "Código", "Descrição", "Entradas", "Saídas", "Saldo",
            "Custo", "Valor Total",
        ]
        ws.append(cabecalho)
        for celula in ws[ws.max_row]:
            celula.font = Font(bold=True)

        for linha in linhas:
            ws.append([
                linha.get("codigo", ""),
                linha.get("descricao", ""),
                round(linha.get("qtd_entradas", 0), 3),
                round(linha.get("qtd_saidas", 0), 3),
                round(linha.get("saldo", 0), 3),
                round(linha.get("custo", 0), 2),
                round(linha.get("valor_total", 0), 2),
            ])

        total_entradas = sum(l.get("qtd_entradas", 0) for l in linhas)
        total_saidas = sum(l.get("qtd_saidas", 0) for l in linhas)
        total_saldo = sum(l.get("saldo", 0) for l in linhas)
        total_valor = sum(l.get("valor_total", 0) for l in linhas)

        ws.append([
            "", "TOTAL",
            round(total_entradas, 3),
            round(total_saidas, 3),
            round(total_saldo, 3),
            "",
            round(total_valor, 2),
        ])
        for celula in ws[ws.max_row]:
            celula.font = Font(bold=True)

        # Larguras de coluna aproximadas
        larguras = [14, 40, 12, 12, 12, 12, 14]
        for i, largura in enumerate(larguras, start=1):
            ws.column_dimensions[ws.cell(
                row=3, column=i
            ).column_letter].width = largura

        wb.save(caminho)

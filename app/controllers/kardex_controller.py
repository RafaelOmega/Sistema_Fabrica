# app/controllers/kardex_controller.py
from PySide6.QtCore import QDate
from PySide6.QtWidgets import QFileDialog, QMessageBox, QWidget

from app.models.kardex_table_model import KardexTableModel
from app.repositories.movimento_estoque_repository import (
    MovimentoEstoqueRepository,
)
from app.repositories.produto_repository import ProdutoRepository
from app.utils.logger import get_logger
from app.utils.table_utils import configurar_tabela
from app.views.ui_movimento_estoque import Ui_Movimento_Estoque

logger = get_logger("kardex_controller")


class KardexController(QWidget):
    def __init__(self):
        super().__init__()
        self.ui = Ui_Movimento_Estoque()
        self.ui.setupUi(self)

        self.repo = MovimentoEstoqueRepository()
        self.produto_repo = ProdutoRepository()
        self.produto_id = None

        self.model = KardexTableModel()

        self._configurar_tabela()
        self._configurar_campos()
        self._conectar_sinais()
        self._estado_inicial()
        logger.info("Tela de kardex inicializada")

    # --- Configuração ---

    def _configurar_tabela(self):
        self.ui.tb_Kardex.setModel(self.model)
        configurar_tabela(self.ui.tb_Kardex, coluna_stretch=1,
                          ordenavel=True)

    def _configurar_campos(self):
        for dt in (self.ui.dt_Data_Inicial, self.ui.dt_Data_Final):
            dt.setCalendarPopup(True)
            dt.setDisplayFormat("dd/MM/yyyy")

        hoje = QDate.currentDate()
        self.ui.dt_Data_Inicial.setDate(QDate(hoje.year(), 1, 1))
        self.ui.dt_Data_Final.setDate(hoje)

    def _conectar_sinais(self):
        self.ui.bt_Pesquisar_Produto.clicked.connect(
            self.abrir_pesquisa_produto)
        self.ui.txt_Cod_Prod.returnPressed.connect(self.buscar_por_codigo)
        self.ui.bt_Atualizar.clicked.connect(self.atualizar_kardex)
        self.ui.bt_Exportar.clicked.connect(self.exportar_excel)

    def _estado_inicial(self):
        self.ui.bt_Atualizar.setEnabled(False)
        self.ui.bt_Exportar.setEnabled(False)

    # --- Seleção de produto ---

    def abrir_pesquisa_produto(self):
        from app.controllers.pesquisa_produto_controller import (
            PesquisaProdutoController,
        )
        dlg = PesquisaProdutoController(self)
        if dlg.exec() and getattr(dlg, "produto_selecionado", None):
            self._definir_produto(dlg.produto_selecionado)

    def buscar_por_codigo(self):
        codigo = self.ui.txt_Cod_Prod.text().strip()
        if not codigo:
            return
        buscar = getattr(self.produto_repo, "buscar_por_codigo", None)
        if buscar is None:
            QMessageBox.information(
                self, "Aviso",
                "Use o botão '...' para localizar o produto.")
            return
        try:
            produto = buscar(codigo)
        except Exception as e:
            logger.error(f"Erro ao buscar produto por código: {e}",
                         exc_info=True)
            QMessageBox.critical(self, "Erro", f"Erro ao buscar: {e}")
            return
        if produto is None:
            QMessageBox.warning(self, "Aviso",
                                f"Produto '{codigo}' não encontrado.")
            return
        self._definir_produto(produto)

    def _definir_produto(self, produto):
        self.produto_id = produto.id
        self.ui.txt_Cod_Prod.setText(str(produto.codigo))
        self.ui.txt_Descricao_Produto.setText(str(produto.descricao))
        self.ui.bt_Atualizar.setEnabled(True)
        self.atualizar_kardex()

    # --- Consulta ---

    def _periodo(self):
        data_inicial = self.ui.dt_Data_Inicial.date().toPython()
        data_final = self.ui.dt_Data_Final.date().toPython()
        if data_inicial > data_final:
            data_inicial, data_final = data_final, data_inicial
            self.ui.dt_Data_Inicial.setDate(data_inicial)
            self.ui.dt_Data_Final.setDate(data_final)
        return data_inicial, data_final

    def atualizar_kardex(self):
        if self.produto_id is None:
            QMessageBox.warning(
                self, "Aviso", "Selecione um produto primeiro.")
            return
        try:
            data_inicial, data_final = self._periodo()
            linhas = self.repo.listar_por_produto(
                self.produto_id, data_inicial, data_final)
            self.model.atualizar_dados(linhas)
            self.ui.bt_Exportar.setEnabled(bool(linhas))
            logger.info(
                f"Kardex atualizado | produto_id={self.produto_id} "
                f"| periodo={data_inicial} a {data_final} "
                f"| movimentos={len(linhas)}")
        except Exception as e:
            logger.error(f"Erro ao atualizar kardex: {e}", exc_info=True)
            QMessageBox.critical(self, "Erro", f"Erro ao atualizar: {e}")

    # --- Exportação ---

    def exportar_excel(self):
        if not self.model.linhas:
            return
        padrao = f"kardex_{self.ui.txt_Cod_Prod.text().strip() or 'produto'}.xlsx"
        caminho, _ = QFileDialog.getSaveFileName(
            self, "Exportar para Excel", padrao, "Excel (*.xlsx)")
        if not caminho:
            return
        try:
            from openpyxl import Workbook
            from openpyxl.styles import Font

            wb = Workbook()
            ws = wb.active
            ws.title = "Kardex"
            ws.append(self.model.HEADER)
            for cell in ws[1]:
                cell.font = Font(bold=True)

            for linha in range(self.model.rowCount()):
                ws.append([
                    self.model.dados_celula(linha, col)
                    for col in range(self.model.columnCount())
                ])

            # Largura de coluna aproximada pelo conteúdo
            for col in range(1, self.model.columnCount() + 1):
                maior = max(
                    len(str(ws.cell(row=r, column=col).value or ""))
                    for r in range(1, self.model.rowCount() + 2)
                )
                ws.column_dimensions[
                    ws.cell(row=1, column=col).column_letter
                ].width = max(maior + 2, 10)

            wb.save(caminho)
            logger.info(f"Kardex exportado para {caminho}")
            QMessageBox.information(
                self, "Sucesso", f"Exportado com sucesso para:\n{caminho}")
        except Exception as e:
            logger.error(f"Erro ao exportar kardex: {e}", exc_info=True)
            QMessageBox.critical(self, "Erro", f"Erro ao exportar: {e}")
